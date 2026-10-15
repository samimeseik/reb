"""CSV / TSV (stdlib) and XLSX (optional `openpyxl`) exports."""

from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Any

from ..timeparse import infer_dayfirst, parse_timestamp
from . import Extraction, FactDraft


def read_rows(path: Path) -> list[dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ImportError("Reading .xlsx needs the 'excel' extra: pip install reb-agents[excel]") from exc
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        header = [str(h).strip() if h is not None else f"col{i}" for i, h in enumerate(rows[0])]
        return [dict(zip(header, r)) for r in rows[1:] if any(c is not None for c in r)]
    raw = path.read_text(encoding="utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(raw[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel_tab if suffix == ".tsv" else csv.excel
    reader = csv.DictReader(io.StringIO(raw), dialect=dialect)
    return [{(k or "").strip(): v for k, v in row.items()} for row in reader]


def _rows_as_text(rows: list[dict[str, Any]]) -> str:
    if not rows:
        return ""
    out = io.StringIO()
    writer = csv.DictWriter(out, fieldnames=list(rows[0].keys()), extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        writer.writerow({k: ("" if v is None else v) for k, v in row.items()})
    return out.getvalue()


def extract(path: Path) -> Extraction:
    rows = read_rows(path)
    columns = list(rows[0].keys()) if rows else []
    facts = [
        FactDraft(
            fact=f"Table with {len(rows)} rows and {len(columns)} columns: {', '.join(columns[:20])}.",
            operation="inventory_table",
            normalized={"rows": len(rows), "columns": columns},
        )
    ]
    for col in columns:
        dayfirst, ambiguous = infer_dayfirst(r.get(col) for r in rows)
        values = [parse_timestamp(r.get(col), dayfirst) for r in rows[:500]]
        parsed = [v for v in values if v is not None]
        if rows and len(parsed) >= max(1, int(0.8 * min(len(rows), 500))):
            all_parsed = [p for p in (parse_timestamp(r.get(col), dayfirst) for r in rows) if p]
            start, end = min(all_parsed), max(all_parsed)
            limitations = ["Coverage is what the export contains, not what the business recorded."]
            if ambiguous:
                limitations.append("Day/month order could not be determined; day-first was assumed.")
            facts.append(
                FactDraft(
                    fact=f"Column '{col}' holds dates from {start.date()} to {end.date()}.",
                    operation="date_coverage",
                    location=f"column {col}",
                    normalized={"column": col, "start": start.isoformat(), "end": end.isoformat(), "dayfirst": dayfirst},
                    limitations=limitations,
                )
            )
    return Extraction(provider="spreadsheet", text=_rows_as_text(rows), facts=facts, rows=rows)
